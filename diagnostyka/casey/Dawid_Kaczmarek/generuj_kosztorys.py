"""Kosztorys naprawy na papierze firmowym Ampere Point (ten sam układ co protokoły naprawy).

Korzysta z funkcji generatora protokołów (repair-protocols/src/repair_protocol.py), więc
nagłówek z logo, kolory i kroje są identyczne. Uruchamiać pythonem z venv tego repo:
  ..\\..\\repair-protocols\\.venv\\Scripts\\python -X utf8 generuj_kosztorys.py kosztorys.json
Powstaje kosztorys.docx i kosztorys.pdf (PDF przez Word, jak w protokołach).
"""
from __future__ import annotations

import json
import sys
from decimal import Decimal
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1] / "repair-protocols" / "src"))
import repair_protocol as rp  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Pt  # noqa: E402


def header_footer(doc, data):
    rp._header_footer(doc, {"protocol": {"number": ""}, "service_provider": data["service_provider"]})
    section = doc.sections[0]
    cell = section.header.tables[0].cell(0, 1)
    cell.text = f"{data['document']['title']}  |  {data['document']['issue_date']}"
    rp._format_cell(cell, bold=True, color=rp.NAVY, size=8.1, align=WD_ALIGN_PARAGRAPH.RIGHT)
    footer = section.footer.paragraphs[0]
    for run in footer.runs:
        run.text = ""
    provider = data["service_provider"]
    run = footer.add_run(f"{provider['name']}  •  {provider['address']}  •  {provider['contact']}")
    rp._set_font(run, size=7.2, color=rp.MID_GRAY)


def title(doc, data):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    rp._set_font(paragraph.add_run(data["document"]["title"]), size=17, bold=True, color=rp.NAVY)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(7)
    rp._set_font(subtitle.add_run(f"{data['document']['place']}, {data['document']['issue_date']}"),
                 size=9.2, color=rp.MID_GRAY)


def detail_table(doc, rows):
    """Jak rp._label_detail_table, ale para (None, None) zostaje pustą, niecieniowaną komórką."""
    table = doc.add_table(rows=0, cols=4)
    table.style = "Table Grid"
    rp._set_table_widths(table, [3.0, 6.0, 3.0, 6.0])
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            label = idx in (0, 2)
            if row[idx - (idx % 2)] is None:
                cells[idx].text = ""
                continue
            cells[idx].text = str(value or "—")
            if label:
                rp._set_cell_shading(cells[idx], rp.LIGHT_BLUE)
            rp._format_cell(cells[idx], bold=label, color=rp.NAVY if label else "000000")


def cost_table(doc, data):
    currency = data["currency"]
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    rp._set_table_widths(table, [1.2, 11.6, 2.0, 3.2])
    for idx, text in enumerate(("Lp.", "Pozycja", "Ilość", "Kwota brutto")):
        cell = table.rows[0].cells[idx]
        cell.text = text
        rp._set_cell_shading(cell, rp.LIGHT_BLUE)
        rp._format_cell(cell, bold=True, color=rp.NAVY, size=8.1, align=WD_ALIGN_PARAGRAPH.CENTER)
    total = Decimal("0")
    for number, item in enumerate(data["items"], start=1):
        cells = table.add_row().cells
        values = (f"{number}.", item["name"], item["quantity"], rp._money(item["gross"], currency))
        aligns = (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                  WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT)
        for idx, value in enumerate(values):
            cells[idx].text = value
            rp._format_cell(cells[idx], size=8.8, align=aligns[idx])
        total += rp._decimal(item["gross"])
    cells = table.add_row().cells
    merged = cells[0].merge(cells[2])
    merged.text = "RAZEM DO ZAPŁATY (brutto)"
    rp._set_cell_shading(merged, rp.LIGHT_BLUE)
    rp._format_cell(merged, bold=True, color=rp.NAVY, size=9.2, align=WD_ALIGN_PARAGRAPH.RIGHT)
    cells[3].text = rp._money(total, currency)
    rp._set_cell_shading(cells[3], rp.LIGHT_BLUE)
    rp._format_cell(cells[3], bold=True, color=rp.NAVY, size=9.2, align=WD_ALIGN_PARAGRAPH.RIGHT)


def prepared_by(doc, name):
    # Bez miejsca na ręczny podpis — kosztorys idzie e-mailem (decyzja Dawida, 2026-10-05).
    paragraph = doc.add_paragraph()
    rp._set_font(paragraph.add_run("Sporządził: "), size=9.2, bold=True, color=rp.NAVY)
    rp._set_font(paragraph.add_run(name), size=9.2)


def build(data):
    doc = Document()
    rp._configure_document(doc)
    header_footer(doc, data)
    title(doc, data)

    provider, customer, equipment = data["service_provider"], data["customer"], data["equipment"]
    rp._section_heading(doc, "1", "Strony")
    detail_table(doc, [
        ("Wykonawca", provider["name"], "Zamawiający", customer["name"]),
        ("Adres", provider["address"], "Adres", customer["address"]),
        ("Kontakt", provider["contact"], None, None),
    ])
    rp._section_heading(doc, "2", "Urządzenie")
    detail_table(doc, [
        ("Model", equipment["model"], "Data zakupu", equipment["purchase_date"]),
        ("Nr seryjny", equipment["serial_number"], None, None),
    ])
    rp._section_heading(doc, "3", "Uszkodzenie i zakres naprawy")
    rp._text_block(doc, "Opis uszkodzenia", data["damage"])
    rp._text_block(doc, "Zakres naprawy", data["scope"])
    rp._section_heading(doc, "4", "Koszt naprawy")
    cost_table(doc, data)
    doc.add_paragraph()
    for note in data["notes"]:
        rp._small_paragraph(doc, note, color="000000", size=8.8, space_after=3)
    doc.add_paragraph()
    prepared_by(doc, data["prepared_by"])

    props = doc.core_properties
    props.title = f"{data['document']['title']} — {equipment['model']}"
    props.author = provider["name"]
    return doc


def main(json_path: Path):
    data = json.loads(json_path.read_text(encoding="utf-8"))
    docx_path = json_path.with_suffix(".docx")
    build(data).save(docx_path)
    pdf_path = rp._create_pdf(docx_path)
    print(f"{docx_path}\n{pdf_path}")


if __name__ == "__main__":
    main(Path(sys.argv[1]).resolve())
